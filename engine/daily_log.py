#!/usr/bin/env python3
"""Daily XAUUSD price log.

Runs once a day after the gold close. Keeps the full history in docs/data/gold_daily_history.csv
(never shrinks) and rebuilds docs/data/XAUUSD_daily_log.xlsx from it:
  Summary     — latest price and changes (formulas)
  Daily log   — one row per trading day: prices, changes (formulas), daily indicators, signals
  Monthly     — one row per month (formulas over the Daily log)
  Read me     — what each column means
"""
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
import agent  # noqa: E402  (reuses fetchers and config)
from indicators import atr, ema, rsi  # noqa: E402

OUT = agent.OUT
CSV = OUT / "gold_daily_history.csv"
XLSX = OUT / "XAUUSD_daily_log.xlsx"
LOG = "'Daily log'"
RNG = 100000  # rows covered by lookup ranges


def rcol(c):
    return f"{LOG}!${c}$2:${c}${RNG}"

# ───────────────────────────── data ─────────────────────────────

def fetch_daily(first_run):
    key = os.getenv("TWELVE_DATA_KEY", "").strip()
    notes = []
    df = None
    if key:
        try:
            df = agent.fetch_twelve(agent.CFG["spot_symbol"], "1day", 1100 if first_run else 40, key)
            source = "Spot XAU/USD (Twelve Data)"
        except Exception as e:  # noqa: BLE001
            notes.append(f"Twelve Data failed: {e}")
    if df is None:
        df = agent.fetch_yf(agent.CFG["futures_symbol"], "1d", "5y" if first_run else "3mo")
        source = "COMEX gold futures GC=F (Yahoo)"
    df = df[["open", "high", "low", "close"]].copy()
    try:
        fut = agent.fetch_yf(agent.CFG["futures_symbol"], "1d", "5y" if first_run else "3mo")
        vol = fut["volume"]
        vol.index = vol.index.tz_convert(None).normalize()
        df.index = df.index.tz_convert(None).normalize()
        df["volume"] = vol.reindex(df.index)
    except Exception as e:  # noqa: BLE001
        df.index = df.index.tz_convert(None).normalize()
        df["volume"] = pd.NA
        notes.append(f"Volume unavailable: {e}")
    return df, source, notes


def merge_history(new):
    if CSV.exists():
        old = pd.read_csv(CSV, parse_dates=["date"]).set_index("date")
        allr = pd.concat([old[~old.index.isin(new.index)], new]).sort_index()
    else:
        allr = new.sort_index()
    allr = allr[~allr.index.duplicated(keep="last")]
    allr = allr[allr.index.dayofweek < 5]                      # no weekend stubs
    today = pd.Timestamp(datetime.now(timezone.utc).date())
    allr = allr[allr.index <= today]
    allr.index.name = "date"
    allr.round(4).to_csv(CSV)
    return allr


def signals_by_day():
    p = OUT / "journal.json"
    if not p.exists():
        return {}
    out = {}
    for s in json.loads(p.read_text()).get("signals", []):
        d = pd.Timestamp(s["time_utc"]).tz_convert(None).normalize()
        row = out.setdefault(d, {"n": 0, "w": 0, "l": 0, "r": 0.0})
        row["n"] += 1
        if s.get("status") == "closed":
            row["r"] += float(s.get("r", 0))
            if s.get("r", 0) > 0:
                row["w"] += 1
            else:
                row["l"] += 1
    return out


# ───────────────────────────── workbook ─────────────────────────────

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
BOLD = Font(name=F, size=10, bold=True)
HDR_FONT = Font(name=F, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="26211A")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
GREEN = PatternFill("solid", fgColor="D9EAD3")
RED = PatternFill("solid", fgColor="F4CCCC")
PRICE = "#,##0.00"
CHG = "+#,##0.00;-#,##0.00;0.00"
PCT = "+0.00%;-0.00%;0.00%"


def header(ws, row, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def build(df, source, notes):
    c = df["close"]
    df = df.assign(ema20=ema(c, 20), ema50=ema(c, 50), ema200=ema(c, 200),
                   rsi14=rsi(c, 14), atr14=atr(df, 14))
    warm = pd.Series(range(len(df)), index=df.index)
    for col, n in (("ema20", 20), ("ema50", 50), ("ema200", 200), ("rsi14", 15), ("atr14", 15)):
        df.loc[warm < n - 1, col] = pd.NA                     # no value until enough history
    sig = signals_by_day()

    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    lg = wb.create_sheet("Daily log")
    mo = wb.create_sheet("Monthly")
    rm = wb.create_sheet("Read me")

    # ── Daily log ──
    cols = ["Date", "Day", "Open ($)", "High ($)", "Low ($)", "Close ($)", "Change ($)", "Change (%)",
            "Range ($)", "Direction", "ATR 14 ($)", "RSI 14", "EMA 20 ($)", "EMA 50 ($)", "EMA 200 ($)",
            "Trend vs EMA 200", "Volume (COMEX)", "Signals", "Wins", "Losses", "Net R"]
    widths = [12, 6, 11, 11, 11, 11, 11, 11, 10, 10, 10, 8, 11, 11, 11, 15, 13, 8, 7, 7, 8]
    header(lg, 1, cols, widths)
    for i, (d, r) in enumerate(df.iterrows()):
        x = i + 2
        s = sig.get(d, {"n": 0, "w": 0, "l": 0, "r": 0.0})
        vals = {1: d.to_pydatetime(), 3: r.open, 4: r.high, 5: r.low, 6: r.close,
                11: r.atr14, 12: r.rsi14, 13: r.ema20, 14: r.ema50, 15: r.ema200,
                17: r.volume, 18: s["n"], 19: s["w"], 20: s["l"], 21: round(s["r"], 2)}
        for col, v in vals.items():
            cell = lg.cell(row=x, column=col, value=None if pd.isna(v) else (round(float(v), 2) if col in (3, 4, 5, 6, 11, 12, 13, 14, 15) else v))
            cell.font = BLUE
        lg.cell(row=x, column=2, value=f'=TEXT(A{x},"ddd")')
        if x > 2:
            lg.cell(row=x, column=7, value=f"=F{x}-F{x-1}")
            lg.cell(row=x, column=8, value=f"=IF(F{x-1}=0,0,G{x}/F{x-1})")
        lg.cell(row=x, column=9, value=f"=D{x}-E{x}")
        lg.cell(row=x, column=10, value=f'=IF(F{x}>C{x},"Up",IF(F{x}<C{x},"Down","Flat"))')
        lg.cell(row=x, column=16, value=f'=IF(O{x}="","",IF(F{x}>O{x},"Above","Below"))')
        for col in (2, 7, 8, 9, 10, 16):
            lg.cell(row=x, column=col).font = BLACK
        for col in range(1, 22):
            lg.cell(row=x, column=col).border = THIN
        lg.cell(row=x, column=1).number_format = "dd mmm yyyy"
        for col in (3, 4, 5, 6, 9, 11, 13, 14, 15):
            lg.cell(row=x, column=col).number_format = PRICE
        lg.cell(row=x, column=7).number_format = CHG
        lg.cell(row=x, column=8).number_format = PCT
        lg.cell(row=x, column=12).number_format = "0.0"
        lg.cell(row=x, column=17).number_format = "#,##0"
        lg.cell(row=x, column=21).number_format = CHG
    last = len(df) + 1
    lg.freeze_panes = "B2"
    lg.auto_filter.ref = f"A1:U{last}"
    lg.conditional_formatting.add(f"G2:H{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN))
    lg.conditional_formatting.add(f"G2:H{last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    lg.conditional_formatting.add(f"J2:J{last}", FormulaRule(formula=[f'J2="Up"'], fill=GREEN))
    lg.conditional_formatting.add(f"J2:J{last}", FormulaRule(formula=[f'J2="Down"'], fill=RED))
    lg.conditional_formatting.add(f"U2:U{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN))
    lg.conditional_formatting.add(f"U2:U{last}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))

    # ── Summary ──
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 18
    sm.column_dimensions["C"].width = 52
    sm["A1"] = "XAUUSD daily log — summary"
    sm["A1"].font = Font(name=F, size=14, bold=True)
    sm["A2"] = f"Source: {source}. Updated automatically after each trading day."
    sm["A2"].font = Font(name=F, size=9, color="666666")
    n = "$B$4"
    rows = [
        ("Days logged", f"=COUNTA({LOG}!A:A)-1", "0", "Trading days in the log"),
        ("Latest date", f"=INDEX({LOG}!A:A,{n}+1)", "dd mmm yyyy", ""),
        ("Latest close ($)", f"=INDEX({LOG}!F:F,{n}+1)", PRICE, ""),
        ("Day change ($)", f"=INDEX({LOG}!G:G,{n}+1)", CHG, ""),
        ("Day change (%)", f"=INDEX({LOG}!H:H,{n}+1)", PCT, ""),
        ("1 week change (%)", f"=IFERROR(B6/INDEX({LOG}!F:F,{n}+1-5)-1,\"\")", PCT, "Versus 5 trading days ago"),
        ("1 month change (%)", f"=IFERROR(B6/INDEX({LOG}!F:F,{n}+1-21)-1,\"\")", PCT, "Versus 21 trading days ago"),
        ("3 month change (%)", f"=IFERROR(B6/INDEX({LOG}!F:F,{n}+1-63)-1,\"\")", PCT, "Versus 63 trading days ago"),
        ("Year to date (%)", f"=IFERROR(B6/INDEX({rcol('F')},MATCH(DATE(YEAR(B5),1,1)-1,{rcol('A')},1))-1,\"\")", PCT,
         "Versus the last close of the previous year"),
        ("52 week high ($)", f"=_xlfn.MAXIFS({LOG}!D:D,{LOG}!A:A,\">=\"&(B5-365))", PRICE, ""),
        ("52 week low ($)", f"=_xlfn.MINIFS({LOG}!E:E,{LOG}!A:A,\">=\"&(B5-365))", PRICE, ""),
        ("Below 52 week high (%)", "=IFERROR(B6/B13-1,\"\")", PCT, ""),
        ("Average daily range, 20 days ($)",
         f"=IFERROR(AVERAGE(INDEX({LOG}!I:I,{n}+1-19):INDEX({LOG}!I:I,{n}+1)),\"\")", PRICE, "High minus low"),
        ("ATR 14 ($)", f"=INDEX({LOG}!K:K,{n}+1)", PRICE, "Typical daily move"),
        ("RSI 14", f"=INDEX({LOG}!L:L,{n}+1)", "0.0", "Above 50 = upward momentum"),
        ("Trend vs EMA 200", f"=INDEX({LOG}!P:P,{n}+1)", "@", "Above = long-term uptrend"),
        ("Signals logged", f"=SUM({LOG}!R:R)", "0", "LONG and SHORT signals from the engine"),
        ("Net result (R)", f"=SUM({LOG}!U:U)", CHG, "Closed signals only"),
    ]
    header(sm, 3, ["Measure", "Value", "Note"], [30, 18, 52])
    for i, (lab, f, fmt, note) in enumerate(rows):
        r = 4 + i
        sm.cell(row=r, column=1, value=lab).font = BLACK
        v = sm.cell(row=r, column=2, value=f)
        v.font, v.number_format = BLACK, fmt
        sm.cell(row=r, column=3, value=note).font = Font(name=F, size=9, color="666666")
        for col in (1, 2, 3):
            sm.cell(row=r, column=col).border = THIN
    for addr in ("B7", "B8", "B9", "B10", "B11", "B12", "B21"):
        sm.conditional_formatting.add(addr, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN))
        sm.conditional_formatting.add(addr, CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    if notes:
        sm.cell(row=24, column=1, value="Notes from the last update").font = BOLD
        for i, t in enumerate(notes):
            sm.cell(row=25 + i, column=1, value=t).font = Font(name=F, size=9, color="666666")

    # ── Monthly ──
    mcols = ["Month", "Open ($)", "High ($)", "Low ($)", "Close ($)", "Change (%)", "Avg daily range ($)",
             "Up days", "Down days", "Trading days", "Signals", "Net R"]
    header(mo, 1, mcols, [11, 11, 11, 11, 11, 11, 13, 8, 9, 9, 8, 8])
    months = sorted({pd.Timestamp(d.year, d.month, 1) for d in df.index})
    for i, m in enumerate(months):
        x = i + 2
        a = f"A{x}"
        nxt = f"EDATE({a},1)"
        inrange = f'{LOG}!A:A,">="&{a},{LOG}!A:A,"<"&{nxt}'
        cell = mo.cell(row=x, column=1, value=m.to_pydatetime())
        cell.font, cell.number_format = BLUE, "mmm yyyy"
        f = {
            2: f"=INDEX({rcol('C')},MATCH({a}-1,{rcol('A')},1)+1)" if i > 0 else f"={LOG}!C2",
            3: f"=_xlfn.MAXIFS({LOG}!D:D,{inrange})",
            4: f"=_xlfn.MINIFS({LOG}!E:E,{inrange})",
            5: f"=INDEX({rcol('F')},MATCH({nxt}-1,{rcol('A')},1))",
            6: f"=IFERROR(E{x}/E{x-1}-1,E{x}/B{x}-1)" if i > 0 else f"=E{x}/B{x}-1",
            7: f"=IFERROR(AVERAGEIFS({LOG}!I:I,{inrange}),\"\")",
            8: f'=COUNTIFS({inrange},{LOG}!J:J,"Up")',
            9: f'=COUNTIFS({inrange},{LOG}!J:J,"Down")',
            10: f"=COUNTIFS({inrange})",
            11: f"=SUMIFS({LOG}!R:R,{inrange})",
            12: f"=SUMIFS({LOG}!U:U,{inrange})",
        }
        for col, formula in f.items():
            c2 = mo.cell(row=x, column=col, value=formula)
            c2.font = BLACK
            c2.number_format = {6: PCT, 8: "0", 9: "0", 10: "0", 11: "0", 12: CHG}.get(col, PRICE)
        for col in range(1, 13):
            mo.cell(row=x, column=col).border = THIN
    ml = len(months) + 1
    mo.freeze_panes = "B2"
    mo.conditional_formatting.add(f"F2:F{ml}", CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN))
    mo.conditional_formatting.add(f"F2:F{ml}", CellIsRule(operator="lessThan", formula=["0"], fill=RED))

    # ── Read me ──
    rm.column_dimensions["A"].width = 22
    rm.column_dimensions["B"].width = 90
    info = [
        ("What this is", "A daily record of gold (XAUUSD) prices, updated automatically after each trading day."),
        ("Price source", source),
        ("Volume", "COMEX gold futures daily volume (Yahoo). Blank when unavailable."),
        ("Signals, Wins, Losses, Net R", "From the gold-agent journal: LONG/SHORT signals that fired that day (UTC date)."),
        ("Blue numbers", "Imported data: prices, indicators, volume, signals."),
        ("Black cells", "Formulas: changes, range, direction, trend, and everything on Summary and Monthly."),
        ("Indicators", "EMA 20/50/200, RSI 14 and ATR 14 on daily candles. Blank until enough history exists."),
        ("Download", "Latest copy: basemayman-ux.github.io/gold-agent/data/XAUUSD_daily_log.xlsx"),
        ("Your own notes", "This file is rebuilt every day. Save a copy before adding your own notes."),
    ]
    for i, (k, v) in enumerate(info, 1):
        rm.cell(row=i, column=1, value=k).font = BOLD
        c3 = rm.cell(row=i, column=2, value=v)
        c3.font, c3.alignment = BLACK, Alignment(wrap_text=True, vertical="top")

    wb.save(XLSX)
    return len(df)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    first = not CSV.exists()
    new, source, notes = fetch_daily(first)
    hist = merge_history(new)
    rows = build(hist, source, notes)
    print(f"Daily log: {rows} trading days, last {hist.index[-1]:%Y-%m-%d} close {hist['close'].iloc[-1]:.2f}")


if __name__ == "__main__":
    main()
